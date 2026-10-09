"""Audit and synchronize Section strengths from an authoritative workbook."""

from collections import defaultdict
from datetime import datetime
from pathlib import Path
import sqlite3

from django.core.management.base import BaseCommand, CommandError
from django.db import connection, transaction
from django.db.models import Count, F, Sum
from openpyxl import load_workbook

from academics.models import Section, Semester
from institutions.models import Program
from scheduling.models import GenerationRun, ScheduleEntry, TimetableVersion
from academics.models import CourseOffering


def alias_identity(year, canonical_name):
    """Return only the explicitly approved historical alias for this identity."""
    key = normalized(canonical_name)
    if year == 2 and key in {'cs2fle', 'csai2cle'}:
        return key.removesuffix('le')
    if year == 3 and key.startswith('csai3') and key[-1] in 'abcdefgh' and len(key) == 6:
        return f'cseai3{key[-1]}'
    return None


def duplicate_target_identity(year, stale_name):
    key = normalized(stale_name)
    if year == 3 and key == 'cseccml3a':
        return 'ccml3a'
    if year == 3 and key == 'cseiotbc3a':
        return 'iotbc3a'
    if year == 4 and key.startswith('cseai4') and len(key) == 7 and key[-1] in 'abcdef':
        return f'csai4{key[-1]}'
    if year == 4 and key == 'cseccml4a':
        return 'ccml4'
    if year == 4 and key == 'cseiotbc4a':
        return 'iotbc4'
    return None


def normalized(value):
    return ''.join(ch for ch in str(value or '').casefold() if ch.isalnum())


def integer(value, field, row_number):
    try:
        number = float(str(value).strip())
        if not number.is_integer() or number <= 0:
            raise ValueError
        return int(number)
    except (TypeError, ValueError, OverflowError):
        raise CommandError(f'Row {row_number}: {field} must be a positive whole number.')


def load_rows(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        headings = [str(value or '').strip().casefold().replace('-', '_').replace(' ', '_') for value in next(rows)]
        required = {'program', 'semester', 'year', 'section', 'total_students'}
        if not required.issubset(headings):
            raise CommandError(f'Missing workbook columns: {", ".join(sorted(required - set(headings)))}')
        result = []
        for number, values in enumerate(rows, 2):
            if all(value is None or str(value).strip() == '' for value in values):
                continue
            row = dict(zip(headings, values))
            result.append((number, row))
        return result
    finally:
        book.close()


def reverse_counts(section):
    counts = {}
    for relation in Section._meta.get_fields():
        if relation.auto_created and not relation.concrete and (relation.one_to_many or relation.one_to_one or relation.many_to_many):
            manager = getattr(section, relation.get_accessor_name())
            count = manager.count() if hasattr(manager, 'count') else int(manager is not None)
            if count:
                counts[relation.related_model._meta.label] = count
    return counts


def duplicate_conflicts(stale, target):
    """Conservatively audit relationships; never assume overlapping schedules merge."""
    conflicts = []
    for accessor, label, fields in (
        ('weekly_off_policies', 'SectionWeeklyOffPolicy', ('academic_session_id', 'semester_id')),
        ('period_delivery_policies', 'SectionDeliveryPolicy', ('academic_session_id', 'semester_id')),
    ):
        old_rows = list(getattr(stale, accessor).all())
        new_rows = list(getattr(target, accessor).all())
        for old in old_rows:
            if any(all(getattr(old, field) == getattr(new, field) for field in fields) for new in new_rows):
                conflicts.append(f'{label} overlaps in the same academic period')
            if old.semester_id != target.semester_id or old.academic_session_id != target.semester.session_id:
                conflicts.append(f'{label} has inconsistent academic period on {old.pk}')
    for offering in stale.course_offerings.all():
        if offering.semester_id != target.semester_id:
            conflicts.append(f'CourseOffering {offering.pk} belongs to a different semester')
    target_offerings = set(target.course_offerings.values_list('semester_id', 'course_id'))
    for semester_id, course_id in stale.course_offerings.values_list('semester_id', 'course_id'):
        if (semester_id, course_id) in target_offerings:
            conflicts.append(f'CourseOffering overlaps semester={semester_id} course={course_id}')
    target_entries = list(target.schedule_entries.select_related('start_slot'))
    stale_offering_ids = set(stale.course_offerings.values_list('pk', flat=True))
    for old in stale.schedule_entries.select_related('start_slot'):
        if old.course_offering_id not in stale_offering_ids:
            conflicts.append(f'ScheduleEntry {old.pk} references an offering outside stale Section')
        for new in target_entries:
            if old.version_id != new.version_id or old.weekday != new.weekday:
                continue
            if old.start_slot.template_id != new.start_slot.template_id:
                continue
            old_range = range(old.start_slot.order, old.start_slot.order + old.block_length)
            new_range = range(new.start_slot.order, new.start_slot.order + new.block_length)
            if set(old_range).intersection(new_range):
                conflicts.append(f'ScheduleEntry block collision version={old.version_id} weekday={old.weekday} entries={old.pk},{new.pk}')
    known = {'academics.SectionWeeklyOffPolicy', 'academics.SectionDeliveryPolicy', 'academics.CourseOffering', 'scheduling.ScheduleEntry'}
    for model_label, count in reverse_counts(stale).items():
        if model_label not in known:
            conflicts.append(f'Unreviewed reverse relation {model_label}: {count}')
    return conflicts


def integrity_counts():
    return {
        'entries': ScheduleEntry.objects.count(),
        'locked_entries': ScheduleEntry.objects.filter(locked=True).count(),
        'offerings': CourseOffering.objects.count(),
        'versions': TimetableVersion.objects.count(),
        'generation_runs': GenerationRun.objects.count(),
    }


def merge_pair(stale, target):
    """Move every reviewed reverse FK, preserving related primary keys."""
    if duplicate_conflicts(stale, target):
        raise CommandError(f'Merge conflicts appeared for {stale.pk}')
    for accessor in ('weekly_off_policies', 'period_delivery_policies', 'course_offerings', 'schedule_entries'):
        for related in list(getattr(stale, accessor).all()):
            related.section = target
            related.save(update_fields=['section'])
    remaining = reverse_counts(stale)
    if remaining:
        raise CommandError(f'Stale Section {stale.pk} still has references: {remaining}')
    stale.delete()


class Command(BaseCommand):
    help = 'Audit or safely sync Section strengths from an Excel workbook.'

    def add_arguments(self, parser):
        parser.add_argument('workbook', type=Path)
        mode = parser.add_mutually_exclusive_group(required=True)
        mode.add_argument('--dry-run', action='store_true')
        mode.add_argument('--apply', action='store_true')

    def handle(self, *args, **options):
        workbook = options['workbook']
        if not workbook.is_file():
            raise CommandError(f'Authoritative workbook not found: {workbook}')
        source = load_rows(workbook)
        if not source:
            raise CommandError('The workbook contains no section rows.')
        programs = list(Program.objects.select_related('department'))
        semesters = list(Semester.objects.select_related('session'))
        existing = list(Section.objects.select_related('program', 'semester__session'))
        by_identity = defaultdict(list)
        for section in existing:
            by_identity[(section.program_id, section.semester_id, section.year, normalized(section.name))].append(section)
        planned = []
        used_ids = set()
        source_keys = set()
        ambiguous = []
        for number, row in source:
            program_text = str(row['program'] or '').strip().casefold()
            program_matches = [p for p in programs if program_text in (p.code.casefold(), p.name.casefold())]
            if len(program_matches) != 1:
                ambiguous.append(f'Row {number}: program {row["program"]!r} matched {len(program_matches)} records')
                continue
            program = program_matches[0]
            year = integer(row['year'], 'Year', number)
            strength = integer(row['total_students'], 'Total Students', number)
            name = str(row['section'] or '').strip()
            if not name:
                raise CommandError(f'Row {number}: Section is blank.')
            semester_text = str(row['semester'] or '').strip().casefold()
            semester_matches = [s for s in semesters if s.session.institution_id == program.department.institution_id and (s.name.casefold() == semester_text or str(s.number) == semester_text)]
            # A workbook without a session column cannot safely select between sessions.
            session_text = str(row.get('session') or row.get('academic_session') or '').strip().casefold()
            if session_text:
                semester_matches = [s for s in semester_matches if s.session.name.casefold() == session_text]
            if len(semester_matches) != 1:
                ambiguous.append(f'Row {number}: semester {row["semester"]!r} matched {len(semester_matches)} sessions')
                continue
            semester = semester_matches[0]
            key = (program.pk, semester.pk, year, normalized(name))
            if key in source_keys:
                ambiguous.append(f'Row {number}: duplicate logical section {name!r}')
                continue
            source_keys.add(key)
            matches = by_identity.get(key, [])
            alias = False
            if not matches:
                historical = alias_identity(year, name)
                if historical:
                    matches = by_identity.get((program.pk, semester.pk, year, historical), [])
                    alias = bool(matches)
            if len(matches) > 1:
                ambiguous.append(f'Row {number}: {name!r} matched {len(matches)} database sections')
                continue
            section = matches[0] if matches else None
            if section:
                used_ids.add(section.pk)
            planned.append((number, program, semester, year, name, strength, section, alias))
        stale = [section for section in existing if section.pk not in used_ids]
        referenced = [(section, reverse_counts(section)) for section in stale]
        matched = sum(section is not None for *_, section, _ in planned)
        mismatches = [(section, strength) for _, _, _, _, _, strength, section, _ in planned if section and section.student_strength != strength]
        missing = sum(section is None for *_, section, _ in planned)
        aliases = [(section, name) for _, _, _, _, name, _, section, alias in planned if alias]
        duplicate_pairs = []
        for stale_section in stale:
            target_name = duplicate_target_identity(stale_section.year, stale_section.name)
            targets = by_identity.get((stale_section.program_id, stale_section.semester_id, stale_section.year, target_name), []) if target_name else []
            if len(targets) == 1 and targets[0].pk in used_ids:
                duplicate_pairs.append((stale_section, targets[0], duplicate_conflicts(stale_section, targets[0])))
        self.stdout.write(f'Authoritative rows: {len(source)}; current DB rows: {len(existing)}; matched: {matched}; strength mismatches: {len(mismatches)}; missing: {missing}; stale/extra: {len(stale)}; ambiguous: {len(ambiguous)}; alias-renames: {len(aliases)}; duplicate candidates: {len(duplicate_pairs)}')
        for section, name in aliases:
            self.stdout.write(f'ALIAS_RENAME {section.pk} {section.name} -> {name}')
        for section, strength in mismatches:
            self.stdout.write(f'STRENGTH_MISMATCH {section.pk} {section.program.code} {section.semester.session.name} {section.semester.name} Year {section.year} {section.name}: {section.student_strength} -> {strength}')
        for section, refs in referenced:
            self.stdout.write(f'STALE_EXTRA {section.pk} {section.program.code} {section.semester.session.name} {section.semester.name} Year {section.year} {section.name} strength={section.student_strength} references={refs}')
        for stale_section, target, conflicts in duplicate_pairs:
            self.stdout.write(f'DUPLICATE_PAIR STALE={stale_section.pk} {stale_section.name} TARGET={target.pk} {target.name} REFERENCES={reverse_counts(stale_section)} CONFLICTS={conflicts} SAFE_TO_MERGE={"NO" if conflicts else "YES"}')
        for issue in ambiguous:
            self.stderr.write(f'AMBIGUOUS {issue}')
        if not options['apply']:
            return
        unpaired = {section.pk for section in stale} - {section.pk for section, _, _ in duplicate_pairs}
        conflicts = [(section.pk, issues) for section, _, issues in duplicate_pairs if issues]
        if ambiguous or missing or unpaired or conflicts:
            raise CommandError(f'Apply stopped: ambiguous={len(ambiguous)} missing={missing} unpaired={len(unpaired)} conflicting_pairs={len(conflicts)}.')
        if connection.vendor != 'sqlite':
            raise CommandError('Apply requires a verified database backup; this command currently supports SQLite only.')
        database_name = str(connection.settings_dict['NAME'])
        if database_name.startswith('file:memorydb_') or database_name == ':memory:':
            self.stdout.write('Backup: skipped for isolated in-memory test database')
        else:
            database = Path(database_name)
            backup = database.with_name(f'{database.stem}_before_section_sync_{datetime.now():%Y%m%d_%H%M%S}.sqlite3')
            with sqlite3.connect(str(database)) as source_db, sqlite3.connect(str(backup)) as target_db:
                source_db.backup(target_db)
            self.stdout.write(f'Backup: {backup}')
        before = integrity_counts()
        sections_before = Section.objects.count()
        with transaction.atomic():
            for _, program, semester, year, name, strength, section, _ in planned:
                if section:
                    if section.student_strength != strength or section.name != name:
                        section.student_strength = strength
                        section.name = name
                        section.save(update_fields=['student_strength', 'name'])
                else:
                    created = Section.objects.create(program=program, semester=semester, year=year, name=name, student_strength=strength)
                    self.stdout.write(f'CREATED {created.pk} {name}')
            for stale_section, target, _ in duplicate_pairs:
                merge_pair(stale_section, target)
            after = integrity_counts()
            if after != before:
                raise CommandError(f'Scheduling counts changed: {before} -> {after}')
            expected_sections = sections_before + missing - len(duplicate_pairs)
            if Section.objects.count() != expected_sections:
                raise CommandError('Section count does not match the reviewed merge plan.')
            if ScheduleEntry.objects.exclude(section_id__in=Section.objects.values('pk')).exists():
                raise CommandError('ScheduleEntry has a dangling Section reference.')
            if ScheduleEntry.objects.exclude(course_offering__section_id__in=Section.objects.values('pk')).exists():
                raise CommandError('ScheduleEntry offering has a dangling Section reference.')
            if ScheduleEntry.objects.exclude(section_id=F('course_offering__section_id')).exists():
                raise CommandError('ScheduleEntry and CourseOffering Section references disagree.')
        totals = list(Section.objects.values('year').annotate(sections=Count('pk'), students=Sum('student_strength')).order_by('year'))
        self.stdout.write(f'Final database totals: {totals}')
