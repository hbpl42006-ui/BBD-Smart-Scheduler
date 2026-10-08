from datetime import date
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import Role, User
from academics.models import AcademicSession, Course, CourseOffering, Section, Semester
from faculty.models import CourseOfferingFaculty, Faculty
from institutions.models import Department, Institution, Program
from rooms.models import Room


class MasterDataExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('export@example.com', 'password', role=Role.READ_ONLY_VIEWER)
        cls.institution = Institution.objects.create(name='Export University', code='EXP')
        cls.department = Department.objects.create(institution=cls.institution, name='Computing', code='COMP')
        cls.program = Program.objects.create(department=cls.department, name='Computer Science', code='CS', duration_years=4)
        cls.session = AcademicSession.objects.create(institution=cls.institution, name='2026-27', start_date=date(2026, 7, 1), end_date=date(2027, 6, 30))
        cls.semester = Semester.objects.create(session=cls.session, name='Semester 1', number=1, type='ODD', start_date=date(2026, 7, 1), end_date=date(2026, 12, 31))
        cls.section = Section.objects.create(program=cls.program, semester=cls.semester, year=1, name='CS-1A', student_strength=42)
        cls.active_course = Course.objects.create(code='EXP101', name='Exportable Course', short_code='EC', credit=3)
        cls.archived_course = Course.objects.create(code='OLD101', name='Archived Course', short_code='OC', credit=2, active=False)
        cls.offering = CourseOffering.objects.create(semester=cls.semester, section=cls.section, course=cls.active_course, weekly_periods=3, default_class_type='LECTURE')
        cls.faculty = Faculty.objects.create(name='A. Faculty', employee_code='FAC001', initials='AF', email='faculty@example.com', department=cls.department)
        CourseOfferingFaculty.objects.create(course_offering=cls.offering, faculty=cls.faculty)
        cls.room = Room.objects.create(code='ONL', building='Online', floor='N/A', capacity=0, room_type='CLASSROOM')

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def _export(self, endpoint, **params):
        basename = {
            'faculty': 'faculty', 'sections': 'section', 'rooms': 'room',
            'courses': 'course', 'course-offerings': 'courseoffering',
        }[endpoint]
        return self.client.get(reverse(f'{basename}-export'), params)

    def test_csv_exports_complete_filtered_queryset_and_escapes_formulas(self):
        Course.objects.create(code='=2+2', name='Formula Safety', short_code='FS', credit=1)
        response = self._export('courses', format='csv', status='all', search='Formula')
        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        self.assertIn('attachment; filename="Courses_', response['Content-Disposition'])
        self.assertTrue(response.content.startswith(b'\xef\xbb\xbf'))
        self.assertIn(b"'=2+2", response.content)
        self.assertNotIn(b'Archived Course', response.content)

    def test_xlsx_is_formatted_and_course_default_status_filter_matches_list(self):
        response = self._export('courses')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=False)
        worksheet = workbook.active
        self.assertEqual(worksheet.freeze_panes, 'A2')
        self.assertEqual(worksheet.auto_filter.ref, worksheet.dimensions)
        rows = list(worksheet.values)
        self.assertEqual(rows[0][:3], ('Course Code', 'Course Name', 'Short Code'))
        self.assertTrue(any(row[0] == 'EXP101' for row in rows[1:]))
        self.assertFalse(any(row[0] == 'OLD101' for row in rows[1:]))

    def test_all_five_resources_have_export_actions(self):
        for endpoint in ('faculty', 'sections', 'rooms', 'courses', 'course-offerings'):
            with self.subTest(endpoint=endpoint):
                response = self._export(endpoint, format='csv')
                self.assertEqual(response.status_code, 200)
                self.assertIn('attachment;', response['Content-Disposition'])

    def test_course_offering_export_includes_assigned_faculty_and_online_delivery_policy(self):
        self.section.delivery_policy = Section.DeliveryPolicy.HYBRID
        self.section.offline_weekday = Section.Weekday.MONDAY
        self.section.save(update_fields=['delivery_policy', 'offline_weekday'])
        response = self._export('course-offerings', format='xlsx')
        self.assertEqual(response.status_code, 200, response.content.decode(errors='replace'))
        self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', response['Content-Type'])
        self.assertIn('attachment; filename="Course_Offerings_', response['Content-Disposition'])
        worksheet = load_workbook(BytesIO(response.content), read_only=True).active
        rows = list(worksheet.values)
        values = next(row for row in rows[1:] if row[2] == 'EXP101')
        self.assertIn('FAC001', values)
        self.assertIn('A. Faculty', values)
        self.assertTrue(any('Hybrid' in str(value) for value in values))

    def test_course_offering_csv_export_has_csv_headers(self):
        response = self._export('course-offerings', format='csv')
        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        self.assertIn('attachment; filename="Course_Offerings_', response['Content-Disposition'])

    def test_export_requires_authentication_and_rejects_unknown_format(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(self._export('courses').status_code, 401)
        self.client.force_authenticate(self.user)
        # DRF rejects unknown format overrides during content negotiation.
        self.assertEqual(self._export('courses', format='pdf').status_code, 404)
