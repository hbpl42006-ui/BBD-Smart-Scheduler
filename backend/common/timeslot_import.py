import datetime
import io
from django.db import transaction
from common.models import TimeSlot, TimeSlotTemplate
from institutions.models import Institution

ALIASES = {'template': 'template', 'template name': 'template', 'template_name': 'template', 'label': 'label', 'start time': 'start_time', 'start_time': 'start_time', 'end time': 'end_time', 'end_time': 'end_time', 'order': 'order', 'is break': 'is_break', 'break': 'is_break', 'is_break': 'is_break'}
def _rows(upload):
    raw = upload.read()
    if upload.name.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        values = next((list(sheet.values) for sheet in book.worksheets if sheet.max_row), [])
        if not values: return []
        headers = [str(value).strip() if value is not None else '' for value in values[0]]
        return [dict(zip(headers, row)) for row in values[1:] if any(value is not None and str(value).strip() for value in row)]
    import csv
    return list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
def _bool(value):
    return {'yes': True, 'true': True, '1': True, 'no': False, 'false': False, '0': False}.get(str(value).strip().lower())
def _time(value):
    if isinstance(value, datetime.time): return value
    return datetime.datetime.strptime(str(value).strip(), '%H:%M').time()
def prepare(upload):
    result = {'total': 0, 'create': 0, 'update': 0, 'invalid': 0, 'errors': [], 'rows': []}
    seen = set(); rows = _rows(upload); result['total'] = len(rows)
    for number, source in enumerate(rows, 2):
        row = {ALIASES.get(str(key).strip().lower(), str(key).strip().lower()): ('' if value is None else str(value).strip()) for key, value in source.items()}
        try:
            for field in ('template', 'label', 'start_time', 'end_time', 'order', 'is_break'):
                if not row.get(field): raise ValueError(f'{field.replace("_", " ").title()} is required.')
            start, end = _time(row['start_time']), _time(row['end_time'])
            if end <= start: raise ValueError('End time must be after start time.')
            order = int(row['order'])
            if order <= 0: raise ValueError('Order must be a positive integer.')
            is_break = _bool(row['is_break'])
            if is_break is None: raise ValueError('Is Break must be Yes, No, TRUE, FALSE, 1, or 0.')
            key = (row['template'].lower(), order)
            if key in seen: raise ValueError('Duplicate order within this template.')
            seen.add(key)
            institution = Institution.objects.filter(code='BBDU').first() or Institution.objects.first()
            if not institution: raise ValueError('No institution is configured.')
            template = TimeSlotTemplate.objects.filter(institution=institution, name__iexact=row['template']).first()
            existing = TimeSlot.objects.filter(template=template, order=order).first() if template else None
            prepared = {'template_name': row['template'], 'label': row['label'], 'start_time': start, 'end_time': end, 'order': order, 'is_break': is_break, 'existing_id': str(existing.pk) if existing else None}
            result['update' if existing else 'create'] += 1; result['rows'].append(prepared)
        except Exception as exc:
            result['invalid'] += 1; result['errors'].append({'row': number, 'message': str(exc)})
    return result
@transaction.atomic
def commit(upload):
    result = prepare(upload)
    if result['invalid']: return result
    institution = Institution.objects.filter(code='BBDU').first() or Institution.objects.first()
    templates = {}
    for row in result['rows']:
        template = templates.setdefault(row['template_name'].lower(), TimeSlotTemplate.objects.filter(institution=institution, name__iexact=row['template_name']).first() or TimeSlotTemplate.objects.create(institution=institution, name=row['template_name']))
        values = {key: row[key] for key in ('label', 'start_time', 'end_time', 'order', 'is_break')}
        if row['existing_id']: TimeSlot.objects.filter(pk=row['existing_id']).update(template=template, **values)
        else: TimeSlot.objects.create(template=template, **values)
    return result
