import csv
from datetime import date
from io import BytesIO, StringIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def export_response(queryset, resource, sheet_name, columns, file_stem):
    """Export a fully filtered queryset as a formatted XLSX or CSV response."""
    headers = [label for label, _getter in columns]
    rows = ([getter(item) for _label, getter in columns] for item in queryset.iterator(chunk_size=500))
    filename = f'{file_stem}_{date.today().isoformat()}'
    if resource == 'csv':
        output = StringIO(newline='')
        writer = csv.writer(output, lineterminator='\r\n')
        writer.writerow(headers)
        writer.writerows([_cell_value(value) for value in row] for row in rows)
        response = HttpResponse('\ufeff' + output.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
        return response

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        worksheet.append([_cell_value(value) for value in row])
    worksheet.freeze_panes = 'A2'
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.columns, 1):
        width = min(max(max((len(str(cell.value or '')) for cell in column), default=8) + 2, 12), 48)
        worksheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def _cell_value(value):
    if value is None:
        return ''
    if isinstance(value, str):
        return f"'{value}" if value.startswith(('=', '+', '-', '@', '\t', '\r')) else value
    if isinstance(value, (int, float, bool, date)):
        return value
    return str(value)
